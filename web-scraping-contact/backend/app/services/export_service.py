import os
import json
from datetime import datetime
from typing import List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.schema import Company, CompanyContact, ContactExport

class ExportService:
    @staticmethod
    def generate_contacts_excel(
        db: Session,
        companies: List[Company],
        batch_id: Optional[int] = None,
        search_history_id: Optional[int] = None,
        custom_filename: Optional[str] = None
    ) -> ContactExport:
        """
        Generate a professional, polished multi-sheet Excel workbook.
        Sheet 1: Company Overview
        Sheet 2: All Contacts (with evidence text & verification status)
        Sheet 3: Search Audit (candidate domains, pages crawled, failure reason)
        """
        wb = openpyxl.Workbook()
        
        # Styles
        navy_header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        blue_header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        slate_header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        regular_font = Font(name="Calibri", size=10)
        link_font = Font(name="Calibri", size=10, color="2563EB", underline="single")
        
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        # -------------------------------------------------------------
        # SHEET 1: Company Overview
        # -------------------------------------------------------------
        ws_overview = wb.active
        ws_overview.title = "Company Overview"
        ws_overview.views.sheetView[0].showGridLines = True

        overview_headers = [
            "Company ID", "Input Name", "Official Name", "Website", "Domain",
            "Industry", "Primary Address", "Primary Email", "Primary Phone", "WhatsApp",
            "Search Status", "Data Quality Summary", "Confidence", "Total Contacts", "Last Updated"
        ]
        ws_overview.append(overview_headers)

        for col_idx in range(1, len(overview_headers) + 1):
            cell = ws_overview.cell(row=1, column=col_idx)
            cell.fill = navy_header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)

        total_contacts_count = 0
        all_contacts_rows = []
        audit_rows = []

        for row_idx, comp in enumerate(companies, start=2):
            primary_email = ""
            primary_phone = ""
            whatsapp_num = ""

            verified_contacts_count = 0
            unverified_contacts_count = 0

            for c in comp.contacts:
                if c.email and not primary_email:
                    primary_email = c.email
                if c.phone and not primary_phone and c.contact_type != "whatsapp":
                    primary_phone = c.phone
                if c.contact_type == "whatsapp" or (c.phone and "whatsapp" in str(c.notes).lower()):
                    if not whatsapp_num:
                        whatsapp_num = c.phone

                if c.verification_status in ["verified_official_website", "official_published", "found_on_public_source"]:
                    verified_contacts_count += 1
                else:
                    unverified_contacts_count += 1

                # Collect row for Sheet 2: All Contacts
                contact_val = c.email or c.phone or c.address or c.url or c.raw_value or ""
                all_contacts_rows.append([
                    comp.input_name,
                    comp.official_name or comp.input_name,
                    c.contact_type.replace("_", " ").title(),
                    c.contact_person_name or "-",
                    c.job_title or "-",
                    contact_val,
                    c.source_url or comp.website or "-",
                    c.source_domain or "-",
                    (c.source_section or "-").replace("_", " ").title(),
                    c.evidence_text or "-",
                    (c.extraction_method or "http_html").replace("_", " ").title(),
                    (c.verification_status or "unverified").replace("_", " ").title(),
                    f"{int((c.confidence_score or 0.8) * 100)}%",
                    c.created_at.strftime("%Y-%m-%d %H:%M") if c.created_at else "-"
                ])
                total_contacts_count += 1

            # Parse audit metadata for Sheet 3
            meta = {}
            if comp.audit_metadata:
                try:
                    meta = json.loads(comp.audit_metadata)
                except Exception:
                    pass

            c_domains = ", ".join([d.get("domain", "") for d in meta.get("candidate_domains_checked", []) if d.get("domain")]) or "-"
            reasons = "; ".join(meta.get("rejections", [])) if meta.get("rejections") else "-"
            
            audit_rows.append([
                comp.input_name,
                comp.official_name or comp.input_name,
                c_domains,
                comp.website or "None",
                meta.get("website_selection_reason") or "Direct Match",
                meta.get("pages_crawled_count", 0),
                "Yes" if meta.get("browser_crawler_used") else "No",
                len(comp.contacts),
                verified_contacts_count,
                unverified_contacts_count,
                comp.search_status or ("COMPLETED" if len(comp.contacts) > 0 else "PARTIAL"),
                reasons,
                meta.get("duration_seconds", 0),
                comp.updated_at.strftime("%Y-%m-%d %H:%M") if comp.updated_at else "-"
            ])

            # Quality summary
            quality_summary = f"{verified_contacts_count} Verified Contacts" if verified_contacts_count > 0 else (
                "Official Website Only (No Direct Contacts)" if comp.website else "No Grounded Source Found"
            )

            row_data = [
                comp.id,
                comp.input_name,
                comp.official_name or comp.input_name,
                comp.website or "Not Found",
                comp.domain or "-",
                comp.industry or "-",
                comp.address or "-",
                primary_email or "Not Found",
                primary_phone or "Not Found",
                whatsapp_num or "Not Found",
                comp.search_status or ("COMPLETED" if len(comp.contacts) > 0 else "PARTIAL"),
                quality_summary,
                f"{int((comp.matching_confidence or 0.5) * 100)}%",
                len(comp.contacts),
                comp.updated_at.strftime("%Y-%m-%d %H:%M") if comp.updated_at else "-"
            ]
            ws_overview.append(row_data)

            # Apply row formatting
            for col_idx in range(1, len(overview_headers) + 1):
                cell = ws_overview.cell(row=row_idx, column=col_idx)
                cell.font = regular_font
                cell.border = thin_border
                if row_idx % 2 == 0:
                    cell.fill = zebra_fill
                
                # Format phone numbers as text
                if col_idx in [9, 10]:
                    cell.number_format = '@'

                # Clickable website
                if col_idx == 4 and comp.website and comp.website.startswith("http"):
                    cell.hyperlink = comp.website
                    cell.font = link_font

        ws_overview.freeze_panes = "A2"
        ws_overview.auto_filter.ref = ws_overview.dimensions

        for col in ws_overview.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws_overview.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

        # -------------------------------------------------------------
        # SHEET 2: All Contacts (with Evidence)
        # -------------------------------------------------------------
        ws_details = wb.create_sheet(title="All Contacts")
        ws_details.views.sheetView[0].showGridLines = True

        detail_headers = [
            "Company Input Name", "Official Company Name", "Contact Type",
            "Contact Person", "Position", "Contact Value / Email / Phone / Address",
            "Source URL", "Source Domain", "Source Section", "Evidence Text",
            "Extraction Method", "Verification Status", "Confidence", "Last Checked"
        ]
        ws_details.append(detail_headers)

        for col_idx in range(1, len(detail_headers) + 1):
            cell = ws_details.cell(row=1, column=col_idx)
            cell.fill = blue_header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for r_idx, c_row in enumerate(all_contacts_rows, start=2):
            ws_details.append(c_row)
            for c_idx in range(1, len(detail_headers) + 1):
                cell = ws_details.cell(row=r_idx, column=c_idx)
                cell.font = regular_font
                cell.border = thin_border
                if r_idx % 2 == 0:
                    cell.fill = zebra_fill
                
                # Format contact value as text
                if c_idx == 6:
                    cell.number_format = '@'

                # Clickable source URL
                if c_idx == 7 and str(cell.value).startswith("http"):
                    cell.hyperlink = str(cell.value)
                    cell.font = link_font

        ws_details.freeze_panes = "A2"
        ws_details.auto_filter.ref = ws_details.dimensions

        for col in ws_details.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws_details.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        # -------------------------------------------------------------
        # SHEET 3: Search Audit
        # -------------------------------------------------------------
        ws_audit = wb.create_sheet(title="Search Audit")
        ws_audit.views.sheetView[0].showGridLines = True

        audit_headers = [
            "Input Company Name", "Official Legal Name", "Candidate Domains Checked",
            "Selected Official Website", "Selection Reason", "Pages Crawled",
            "Browser Used", "Contacts Found", "Verified Contacts", "Unverified Contacts",
            "Search Status", "Quality Gate / Rejection Notes", "Duration (s)", "Timestamp"
        ]
        ws_audit.append(audit_headers)

        for col_idx in range(1, len(audit_headers) + 1):
            cell = ws_audit.cell(row=1, column=col_idx)
            cell.fill = slate_header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for r_idx, a_row in enumerate(audit_rows, start=2):
            ws_audit.append(a_row)
            for c_idx in range(1, len(audit_headers) + 1):
                cell = ws_audit.cell(row=r_idx, column=c_idx)
                cell.font = regular_font
                cell.border = thin_border
                if r_idx % 2 == 0:
                    cell.fill = zebra_fill

                # Clickable website
                if c_idx == 4 and str(cell.value).startswith("http"):
                    cell.hyperlink = str(cell.value)
                    cell.font = link_font

        ws_audit.freeze_panes = "A2"
        ws_audit.auto_filter.ref = ws_audit.dimensions

        for col in ws_audit.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws_audit.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        # Save file
        timestamp = datetime.now().strftime("%Y-%m-%d-%H%M%S")
        if custom_filename:
            filename = custom_filename
        else:
            filename = f"company-contacts-{timestamp}.xlsx"

        file_path = str(settings.EXPORT_DIR / filename)
        wb.save(file_path)

        # Record export in database
        export_record = ContactExport(
            batch_id=batch_id,
            search_history_id=search_history_id,
            filename=filename,
            file_path=file_path,
            export_format="xlsx",
            total_companies=len(companies),
            total_contacts=total_contacts_count,
            status="COMPLETED",
            created_at=datetime.utcnow()
        )
        db.add(export_record)
        db.commit()
        db.refresh(export_record)

        return export_record

export_service = ExportService()
