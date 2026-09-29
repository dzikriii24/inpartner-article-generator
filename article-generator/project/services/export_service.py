import io
import re
from datetime import datetime
import markdown as md_lib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

import urllib.request
import tempfile
import uuid
import os

def export_markdown(article) -> str:
    """
    Generates clean Markdown string for the article including metadata and references.
    """
    content = f"# {article.title}\n\n"
    if article.subtitle:
        content += f"> *{article.subtitle}*\n\n"
        
    content += f"**Category:** {article.category or 'Editorial'}  \n"
    content += f"**Author:** {article.author or 'Inpartner Editorial Board'}  \n"
    if article.generated_at:
        content += f"**Date:** {article.generated_at.strftime('%B %d, %Y')}  \n"
    content += "\n---\n\n"
    
    if article.hero_image_url:
        content += f"![{article.title}]({article.hero_image_url})\n"
        if article.hero_image_caption:
            content += f"*{article.hero_image_caption}*\n\n"
            
    content += article.content or ""
    
    if article.sources:
        content += "\n\n## References & Sources\n\n"
        for idx, source in enumerate(article.sources, 1):
            pub = f"**{source.publisher}**" if source.publisher else "Source"
            date_str = source.published_at.strftime('%Y-%m-%d') if source.published_at else ""
            content += f"{idx}. {pub}: [{source.title}]({source.url}) {f'({date_str})' if date_str else ''}\n"
            if source.description:
                clean_desc = re.sub(r'<[^>]+>', '', source.description)[:200]
                content += f"   > {clean_desc}...\n"
                
    return content


def export_html(article) -> str:
    """
    Generates clean HTML document for CMS or offline viewing.
    """
    body_html = md_lib.markdown(article.content or "", extensions=['extra', 'nl2br', 'tables'])
    
    sources_html = ""
    if article.sources:
        sources_html += "<section class='references' style='margin-top: 3rem; padding-top: 2rem; border-top: 2px solid #222;'>\n"
        sources_html += "<h2 style='font-family: sans-serif; font-size: 1.25rem;'>References & Sources</h2>\n"
        sources_html += "<ol style='padding-left: 1.25rem; font-family: sans-serif; font-size: 0.95rem; line-height: 1.6;'>\n"
        for source in article.sources:
            pub = f"<strong>{source.publisher}</strong> &mdash; " if source.publisher else ""
            sources_html += f"<li style='margin-bottom: 0.75rem;'>{pub}<a href='{source.url}' target='_blank' style='color: #93602a; font-weight: 600;'>{source.title}</a>"
            if source.description:
                clean_desc = re.sub(r'<[^>]+>', '', source.description)[:220]
                sources_html += f"<br><span style='color: #666; font-size: 0.88rem;'>{clean_desc}...</span>"
            sources_html += "</li>\n"
        sources_html += "</ol>\n</section>\n"

    full_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{article.title} — Inpartner Article Generator</title>
    <style>
        body {{
            font-family: 'Georgia', 'Times New Roman', serif;
            line-height: 1.8;
            color: #1a1a1a;
            max-width: 800px;
            margin: 0 auto;
            padding: 2rem 1.5rem 4rem;
            background: #ffffff;
        }}
        h1 {{
            font-size: 2.5rem;
            line-height: 1.2;
            color: #111111;
            margin-bottom: 0.5rem;
        }}
        .deck {{
            font-size: 1.3rem;
            color: #555555;
            font-style: italic;
            margin-bottom: 1.5rem;
        }}
        .meta {{
            font-family: sans-serif;
            font-size: 0.85rem;
            color: #777777;
            border-top: 1px solid #eee;
            border-bottom: 1px solid #eee;
            padding: 0.75rem 0;
            margin-bottom: 2rem;
        }}
        .hero-img {{
            width: 100%;
            height: auto;
            border-radius: 6px;
            margin-bottom: 0.5rem;
        }}
        .hero-caption {{
            font-family: sans-serif;
            font-size: 0.82rem;
            color: #777;
            font-style: italic;
            margin-bottom: 2rem;
        }}
        article {{
            font-size: 1.15rem;
        }}
        article h2 {{
            font-size: 1.6rem;
            margin-top: 2.2rem;
            margin-bottom: 0.8rem;
            color: #111;
        }}
        article blockquote {{
            margin: 1.5rem 0;
            padding: 1rem 1.5rem;
            background: #f9f6f0;
            border-left: 4px solid #93602a;
            font-style: italic;
        }}
    </style>
</head>
<body>
    <header>
        <h1>{article.title}</h1>
        {f'<div class="deck">{article.subtitle}</div>' if article.subtitle else ''}
        <div class="meta">
            By <strong>{article.author or "Inpartner Editorial Board"}</strong> &bull;
            {article.category or "Editorial"} &bull;
            {article.generated_at.strftime('%B %d, %Y') if article.generated_at else ''}
        </div>
        {f'<div><img src="{article.hero_image_url}" class="hero-img" alt=""><div class="hero-caption">{article.hero_image_caption or ""}</div></div>' if article.hero_image_url else ''}
    </header>
    <article>
        {body_html}
    </article>
    {sources_html}
</body>
</html>"""
    return full_html


def export_docx(article) -> io.BytesIO:
    """
    Generates formatted Word DOCX document for the article.
    """
    doc = Document()
    
    # Title
    p_title = doc.add_paragraph()
    run_title = p_title.add_run(article.title)
    run_title.font.name = 'Georgia'
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x16, 0x15, 0x13)
    p_title.paragraph_format.space_after = Pt(8)
    
    # Subtitle
    if article.subtitle:
        p_sub = doc.add_paragraph()
        run_sub = p_sub.add_run(article.subtitle)
        run_sub.font.name = 'Georgia'
        run_sub.font.size = Pt(13)
        run_sub.font.italic = True
        run_sub.font.color.rgb = RGBColor(0x4A, 0x47, 0x3F)
        p_sub.paragraph_format.space_after = Pt(14)
        
    # Metadata line
    p_meta = doc.add_paragraph()
    meta_text = f"Category: {article.category or 'Editorial'} | Author: {article.author or 'Inpartner Editorial Board'}"
    if article.generated_at:
        meta_text += f" | Published: {article.generated_at.strftime('%B %d, %Y')}"
    run_meta = p_meta.add_run(meta_text)
    run_meta.font.name = 'Arial'
    run_meta.font.size = Pt(9)
    run_meta.font.color.rgb = RGBColor(0x8C, 0x87, 0x79)
    p_meta.paragraph_format.space_after = Pt(20)
    
    # Body parsing from markdown content
    raw_content = article.content or ""
    lines = raw_content.split('\n')
    
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
            
        if stripped.startswith('# '):
            continue # Already handled title
        elif stripped.startswith('## '):
            heading_text = stripped[3:].strip()
            p_h2 = doc.add_paragraph()
            run_h2 = p_h2.add_run(heading_text)
            run_h2.font.name = 'Georgia'
            run_h2.font.size = Pt(16)
            run_h2.font.bold = True
            run_h2.font.color.rgb = RGBColor(0x93, 0x60, 0x2A)
            p_h2.paragraph_format.space_before = Pt(16)
            p_h2.paragraph_format.space_after = Pt(6)
        elif stripped.startswith('### '):
            heading_text = stripped[4:].strip()
            p_h3 = doc.add_paragraph()
            run_h3 = p_h3.add_run(heading_text)
            run_h3.font.name = 'Georgia'
            run_h3.font.size = Pt(13)
            run_h3.font.bold = True
            p_h3.paragraph_format.space_before = Pt(12)
            p_h3.paragraph_format.space_after = Pt(4)
        elif stripped.startswith('> '):
            quote_text = stripped[2:].strip()
            p_q = doc.add_paragraph()
            p_q.paragraph_format.left_indent = Inches(0.4)
            run_q = p_q.add_run(f'"{quote_text}"')
            run_q.font.name = 'Georgia'
            run_q.font.size = Pt(11)
            run_q.font.italic = True
            p_q.paragraph_format.space_after = Pt(10)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            bullet_text = stripped[2:].strip()
            p_b = doc.add_paragraph(style='List Bullet')
            run_b = p_b.add_run(bullet_text)
            run_b.font.name = 'Georgia'
            run_b.font.size = Pt(11)
            p_b.paragraph_format.space_after = Pt(4)
        else:
            # Clean markdown bold/italic tags if simple
            clean_text = re.sub(r'\*\*(.*?)\*\*', r'\1', stripped)
            # Remove any raw HTML tags that might have been saved by the WYSIWYG editor
            clean_text = re.sub(r'<[^>]+>', '', clean_text)
            
            p_p = doc.add_paragraph()
            run_p = p_p.add_run(clean_text)
            run_p.font.name = 'Georgia'
            run_p.font.size = Pt(11)
            p_p.paragraph_format.space_after = Pt(10)
            
    # References section
    if article.sources:
        p_ref_head = doc.add_paragraph()
        run_rh = p_ref_head.add_run("References & Sources")
        run_rh.font.name = 'Georgia'
        run_rh.font.size = Pt(14)
        run_rh.font.bold = True
        p_ref_head.paragraph_format.space_before = Pt(20)
        p_ref_head.paragraph_format.space_after = Pt(8)
        
        for idx, source in enumerate(article.sources, 1):
            p_s = doc.add_paragraph()
            p_s.paragraph_format.left_indent = Inches(0.2)
            publisher = source.publisher or "Publisher"
            r_idx = p_s.add_run(f"{idx}. [{publisher}] ")
            r_idx.font.bold = True
            r_idx.font.size = Pt(10)
            
            r_title = p_s.add_run(f"{source.title} — {source.url}")
            r_title.font.size = Pt(10)
            p_s.paragraph_format.space_after = Pt(4)

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream


import base64
import html
from PIL import Image
from bs4 import BeautifulSoup, NavigableString

def process_image_for_pdf(src: str) -> str:
    """
    Downloads or decodes an image from any src (URL, base64 data URL, or static path),
    converts it with PIL to a ReportLab-compatible PNG/JPEG, and returns the absolute local path with forward slashes.
    """
    if not src:
        return None
        
    src = html.unescape(src.strip())
    img_bytes = None
    
    try:
        if src.startswith("data:image/"):
            # Handle Base64 Data URL (stored in database)
            header, encoded = src.split(",", 1)
            img_bytes = base64.b64decode(encoded)
        elif src.startswith(("http://", "https://")):
            # Handle Remote Web URL
            req = urllib.request.Request(
                src,
                headers={
                    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'
                }
            )
            with urllib.request.urlopen(req, timeout=12) as response:
                img_bytes = response.read()
        elif src.startswith("/static/"):
            # Handle Local Static File
            base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
            local_path = os.path.join(base_dir, src.lstrip('/'))
            if os.path.exists(local_path):
                with open(local_path, "rb") as f:
                    img_bytes = f.read()
        else:
            # Direct file path check
            if os.path.exists(src):
                with open(src, "rb") as f:
                    img_bytes = f.read()

        if not img_bytes:
            return None

        # Open image with PIL to convert WebP/AVIF/GIF/PNG/etc to RGB JPEG/PNG for ReportLab
        pil_img = Image.open(io.BytesIO(img_bytes))
        out_format = "PNG" if pil_img.mode in ("RGBA", "LA", "P") else "JPEG"
        if out_format == "JPEG" and pil_img.mode != "RGB":
            pil_img = pil_img.convert("RGB")
            
        temp_filename = f"pdf_img_{uuid.uuid4().hex}.{'png' if out_format == 'PNG' else 'jpg'}"
        temp_path = os.path.join(tempfile.gettempdir(), temp_filename)
        pil_img.save(temp_path, format=out_format, quality=92)
        
        # ReportLab XML parser chokes on backslashes in Windows absolute paths
        return temp_path.replace('\\', '/')
    except Exception as e:
        print(f"Error processing image for PDF (src={src[:60]}...): {e}")
        return None

def clean_paragraph_html(p_element):
    """
    Sanitizes HTML inside a paragraph to ensure ReportLab paraparser compatibility.
    """
    # Remove img tags as they are handled as separate flowables
    for img in p_element.find_all('img'):
        img.decompose()
        
    html_str = str(p_element)
    # Strip wrapping <p>, <div>, or <figure> container tags
    html_str = re.sub(r'^\s*<(p|div|figure)[^>]*>', '', html_str, flags=re.IGNORECASE)
    html_str = re.sub(r'</(p|div|figure)>\s*$', '', html_str, flags=re.IGNORECASE)
    
    # Replace em/strong with i/b
    html_str = re.sub(r'<em[^>]*>', '<i>', html_str, flags=re.IGNORECASE)
    html_str = re.sub(r'</em>', '</i>', html_str, flags=re.IGNORECASE)
    html_str = re.sub(r'<strong[^>]*>', '<b>', html_str, flags=re.IGNORECASE)
    html_str = re.sub(r'</strong>', '</b>', html_str, flags=re.IGNORECASE)
    
    # Remove unsupported block/container tags
    html_str = re.sub(r'</?(div|figure|figcaption|span)[^>]*>', '', html_str, flags=re.IGNORECASE)
    html_str = re.sub(r'<br\s*/?>', '<br/>', html_str, flags=re.IGNORECASE)
    
    return html_str.strip()


from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable, KeepTogether, Image as RLImage
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT, TA_JUSTIFY
from reportlab.pdfgen import canvas
import base64
import html
from PIL import Image

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont('Helvetica', 8)
        self.setFillColor(colors.HexColor('#858073'))
        if self._pageNumber > 1:
            self.drawString(45, 805, 'INPARTNER DIGITAL PUBLISHING WORKSPACE')
            self.drawRightString(550, 805, f'Page {self._pageNumber} of {page_count}')
            self.setStrokeColor(colors.HexColor('#E6E1D6'))
            self.setLineWidth(0.5)
            self.line(45, 797, 550, 797)
        else:
            self.drawRightString(550, 30, f'Page 1 of {page_count}')
            self.drawString(45, 30, 'Inpartner Editorial Publication')
            self.setStrokeColor(colors.HexColor('#E6E1D6'))
            self.setLineWidth(0.5)
            self.line(45, 42, 550, 42)
        self.restoreState()


import ssl
import urllib.parse

def process_image_for_pdf(src: str) -> str:
    """
    Downloads or decodes an image from any src (Remote URL, local static URL, base64 data URL, or file path),
    converts it with PIL to a ReportLab-compatible PNG/JPEG, and returns the absolute local path with forward slashes.
    """
    if not src:
        return None
        
    src = html.unescape(str(src).strip())
    img_bytes = None
    
    try:
        # 1. Base64 Data URL
        if src.startswith("data:image/"):
            try:
                header, encoded = src.split(",", 1)
                encoded = urllib.parse.unquote(encoded)
                img_bytes = base64.b64decode(encoded)
            except Exception as e_b64:
                print(f"Base64 decode error: {e_b64}")

        # 2. Local static file or 127.0.0.1 / localhost static URL (prevents FastAPI self-request deadlocks)
        if not img_bytes:
            clean_path = src
            clean_path = re.sub(r'^https?://(127\.0\.0\.1|localhost)(:\d+)?', '', clean_path, flags=re.IGNORECASE)
            
            if clean_path.startswith("/static/") or clean_path.startswith("static/"):
                rel_path = clean_path.lstrip('/')
                base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
                
                candidate_paths = [
                    os.path.join(base_dir, "api", rel_path),
                    os.path.join(base_dir, rel_path),
                    os.path.join(base_dir, "api", "static", rel_path.replace("static/", "")),
                ]
                for c_path in candidate_paths:
                    if os.path.exists(c_path):
                        with open(c_path, "rb") as f:
                            img_bytes = f.read()
                        break

        # 3. Remote HTTP / HTTPS URL
        if not img_bytes and src.startswith(("http://", "https://")):
            try:
                req = urllib.request.Request(
                    src,
                    headers={
                        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'
                    }
                )
                ssl_ctx = ssl.create_default_context()
                ssl_ctx.check_hostname = False
                ssl_ctx.verify_mode = ssl.CERT_NONE
                with urllib.request.urlopen(req, timeout=12, context=ssl_ctx) as response:
                    img_bytes = response.read()
            except Exception as e_url:
                print(f"urllib download error for {src[:60]}: {e_url}")
                try:
                    import requests
                    r = requests.get(
                        src,
                        headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'},
                        timeout=12,
                        verify=False
                    )
                    if r.status_code == 200:
                        img_bytes = r.content
                except Exception as e_req:
                    print(f"requests fallback error for {src[:60]}: {e_req}")

        # 4. Direct local file path check
        if not img_bytes and os.path.exists(src):
            with open(src, "rb") as f:
                img_bytes = f.read()

        if not img_bytes:
            print(f"Failed to acquire image bytes for src: {src[:80]}")
            return None

        # Convert image with PIL to RGB JPEG/PNG for ReportLab compatibility
        pil_img = Image.open(io.BytesIO(img_bytes))
        out_format = "PNG" if pil_img.mode in ("RGBA", "LA", "P") else "JPEG"
        if out_format == "JPEG" and pil_img.mode != "RGB":
            pil_img = pil_img.convert("RGB")
            
        temp_filename = f"pdf_img_{uuid.uuid4().hex}.{'png' if out_format == 'PNG' else 'jpg'}"
        temp_path = os.path.join(tempfile.gettempdir(), temp_filename)
        pil_img.save(temp_path, format=out_format, quality=92)
        
        return temp_path.replace('\\', '/')
    except Exception as e:
        print(f"Error processing image for PDF (src={src[:60]}...): {e}")
        return None

def extract_alignment(elem):
    if not elem:
        return TA_LEFT
    style_attr = elem.get('style', '')
    class_attr = elem.get('class', [])
    if isinstance(class_attr, list):
        class_attr = ' '.join(class_attr)
    align_attr = elem.get('align', '')
    data_align = elem.get('data-align', '')

    combined = f"{style_attr} {class_attr} {align_attr} {data_align}".lower()
    if 'center' in combined:
        return TA_CENTER
    elif 'right' in combined:
        return TA_RIGHT
    elif 'justify' in combined:
        return TA_JUSTIFY
    return TA_LEFT

def clean_paragraph_html(p_element):
    """
    Sanitizes HTML inside a paragraph to ensure ReportLab paraparser compatibility.
    Preserves links <a>, <b>, <i>, <u>, <code> tags.
    """
    for img in p_element.find_all('img'):
        img.decompose()
        
    html_str = str(p_element)
    html_str = re.sub(r'^\s*<(p|div|figure|h[1-6]|blockquote)[^>]*>', '', html_str, flags=re.IGNORECASE)
    html_str = re.sub(r'</(p|div|figure|h[1-6]|blockquote)>\s*$', '', html_str, flags=re.IGNORECASE)
    
    html_str = re.sub(r'<em[^>]*>', '<i>', html_str, flags=re.IGNORECASE)
    html_str = re.sub(r'</em>', '</i>', html_str, flags=re.IGNORECASE)
    html_str = re.sub(r'<strong[^>]*>', '<b>', html_str, flags=re.IGNORECASE)
    html_str = re.sub(r'</strong>', '</b>', html_str, flags=re.IGNORECASE)
    html_str = re.sub(r'</?(span|figure|figcaption)[^>]*>', '', html_str, flags=re.IGNORECASE)
    html_str = re.sub(r'<br\s*/?>', '<br/>', html_str, flags=re.IGNORECASE)
    
    return html_str.strip()


def export_pdf(article) -> io.BytesIO:
    """
    Generates styled PDF document using ReportLab with support for DB base64 images and URL images.
    Supports text alignment, image sizing, captions, credits, page break safety, and hyperlinks.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=45, leftMargin=45,
        topMargin=50, bottomMargin=50,
        title=article.title or "Editorial Article",
        author=article.author or "Inpartner Editorial Board"
    )
    
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'ArticleTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=22,
        leading=26,
        textColor=colors.HexColor('#161513'),
        spaceAfter=10
    )
    
    subtitle_style = ParagraphStyle(
        'ArticleSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#4A473F'),
        spaceAfter=12
    )
    
    meta_style = ParagraphStyle(
        'ArticleMeta',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#8C8779'),
        spaceAfter=15
    )
    
    h1_style = ParagraphStyle(
        'ArticleH1',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#161513'),
        spaceBefore=18,
        spaceAfter=8,
        keepWithNext=True
    )

    h2_style = ParagraphStyle(
        'ArticleH2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#93602A'),
        spaceBefore=16,
        spaceAfter=8,
        keepWithNext=True
    )

    h3_style = ParagraphStyle(
        'ArticleH3',
        parent=styles['Heading3'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=15,
        textColor=colors.HexColor('#161513'),
        spaceBefore=14,
        spaceAfter=6,
        keepWithNext=True
    )
    
    body_style = ParagraphStyle(
        'ArticleBody',
        parent=styles['Normal'],
        fontName='Times-Roman',
        fontSize=11,
        leading=16.5,
        textColor=colors.HexColor('#2C2A26'),
        spaceAfter=10
    )
    
    quote_style = ParagraphStyle(
        'ArticleQuote',
        parent=styles['Normal'],
        fontName='Times-Italic',
        fontSize=11,
        leading=16,
        textColor=colors.HexColor('#161513'),
        leftIndent=20,
        rightIndent=20,
        spaceBefore=8,
        spaceAfter=12
    )

    caption_style = ParagraphStyle(
        'ImageCaption',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor('#656053'),
        alignment=TA_CENTER,
        spaceAfter=4
    )

    credit_style = ParagraphStyle(
        'ImageCredit',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor('#858073'),
        alignment=TA_CENTER,
        spaceAfter=12
    )

    code_style = ParagraphStyle(
        'ArticleCode',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor('#24292e'),
        backColor=colors.HexColor('#f6f8fa'),
        borderColor=colors.HexColor('#e1e4e8'),
        borderWidth=0.5,
        borderPadding=6,
        spaceBefore=8,
        spaceAfter=10
    )
    
    ref_head_style = ParagraphStyle(
        'RefHead',
        parent=styles['Heading3'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=15,
        textColor=colors.HexColor('#161513'),
        spaceBefore=20,
        spaceAfter=10,
        keepWithNext=True
    )
    
    ref_body_style = ParagraphStyle(
        'RefBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#4A473F'),
        leftIndent=12,
        spaceAfter=6
    )

    callout_style = ParagraphStyle(
        'ArticleCallout',
        parent=styles['Normal'],
        fontName='Times-Roman',
        fontSize=10.5,
        leading=15,
        textColor=colors.HexColor('#1D5288'),
        backColor=colors.HexColor('#EEF3F8'),
        borderColor=colors.HexColor('#1D5288'),
        borderWidth=1,
        borderPadding=8,
        spaceBefore=10,
        spaceAfter=12
    )

    def get_style(base, align):
        if align == TA_LEFT:
            return base
        return ParagraphStyle(f"{base.name}_{align}", parent=base, alignment=align)

    def clean_inline(node):
        if not node:
            return ""
        if isinstance(node, str):
            return html.escape(node)
        
        # Create clone to modify
        clone = BeautifulSoup(str(node), 'html.parser')
        root = clone.find() or clone
        for img in root.find_all('img'):
            img.decompose()
            
        inner = root.decode_contents() if hasattr(root, 'decode_contents') else str(root)
        inner = re.sub(r'<strong[^>]*>', '<b>', inner, flags=re.IGNORECASE)
        inner = re.sub(r'</strong>', '</b>', inner, flags=re.IGNORECASE)
        inner = re.sub(r'<em[^>]*>', '<i>', inner, flags=re.IGNORECASE)
        inner = re.sub(r'</em>', '</i>', inner, flags=re.IGNORECASE)
        inner = re.sub(r'<mark[^>]*>', "<font color='#93602A'><b>", inner, flags=re.IGNORECASE)
        inner = re.sub(r'</mark>', "</b></font>", inner, flags=re.IGNORECASE)
        inner = re.sub(r'<code[^>]*>', "<font name='Courier'>", inner, flags=re.IGNORECASE)
        inner = re.sub(r'</code>', "</font>", inner, flags=re.IGNORECASE)
        inner = re.sub(r'<br\s*/?>', '<br/>', inner, flags=re.IGNORECASE)
        inner = re.sub(r'</?(div|p|span|figure|figcaption|h[1-6]|blockquote|ul|ol|li)[^>]*>', '', inner, flags=re.IGNORECASE)
        return inner.strip()
    
    elements = []
    
    # Header branding kicker
    elements.append(Paragraph("<font size=8 color='#93602A'><b>INPARTNER ARTICLE GENERATOR — EDITORIAL REPORT</b></font>", meta_style))
    elements.append(Spacer(1, 4))
    
    # Title & Subtitle
    elements.append(Paragraph(article.title or "Untitled Article", title_style))
    if article.subtitle:
        elements.append(Paragraph(article.subtitle, subtitle_style))
        
    # Metadata Line
    date_str = article.generated_at.strftime('%B %d, %Y') if article.generated_at else datetime.now().strftime('%B %d, %Y')
    meta_str = f"Category: <b>{article.category or 'Editorial'}</b> &bull; Author: <b>{article.author or 'Inpartner Editorial Board'}</b> &bull; Date: <b>{date_str}</b>"
    elements.append(Paragraph(meta_str, meta_style))
    elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#E6E2D8'), spaceAfter=15))
    
    # Hero Image if present
    if article.hero_image_url:
        hero_local = process_image_for_pdf(article.hero_image_url)
        if hero_local:
            hero_block = []
            try:
                pil_img = Image.open(hero_local)
                w, h = pil_img.size
                target_w = 505
                target_h = min(300, int((target_w / w) * h))
                hero_block.append(RLImage(hero_local, width=target_w, height=target_h))
            except Exception:
                hero_block.append(Paragraph(f'<img src="{hero_local}" width="505" />', body_style))
            if article.hero_image_caption:
                hero_block.append(Spacer(1, 4))
                hero_block.append(Paragraph(html.escape(article.hero_image_caption), caption_style))
            hero_block.append(Spacer(1, 10))
            elements.append(KeepTogether(hero_block))
                
    # Content parsing (handles both HTML from WYSIWYG editor and raw Markdown)
    raw_content = article.content or ""
    if not any(tag in raw_content for tag in ['<p', '<h', '<div', '<img', '<figure', '<blockquote']):
        raw_content = md_lib.markdown(raw_content, extensions=['extra', 'nl2br', 'tables'])

    soup = BeautifulSoup(raw_content, 'html.parser')
    root_node = soup.body if soup.body else soup

    # Decompose duplicate H1 or subtitle if present at top of body content
    first_h1 = root_node.find('h1')
    if first_h1 and article.title:
        h1_txt = first_h1.get_text(strip=True).lower()
        art_txt = article.title.strip().lower()
        if h1_txt == art_txt or art_txt.startswith(h1_txt[:25]) or h1_txt.startswith(art_txt[:25]):
            first_h1.decompose()

    if article.subtitle:
        sub_clean = article.subtitle.strip().lower()
        for p in root_node.find_all(['p', 'div'], limit=3):
            p_txt = p.get_text(strip=True).lower()
            if p_txt == sub_clean or sub_clean.startswith(p_txt[:25]) or p_txt.startswith(sub_clean[:25]):
                p.decompose()
                break

    def process_node(node):
        if isinstance(node, NavigableString):
            text = str(node).strip()
            if text:
                elements.append(Paragraph(html.escape(text), body_style))
            return

        if not hasattr(node, 'name') or not node.name:
            return

        align_enum = extract_alignment(node)
        node_classes_str = ' '.join(node.get('class', [])) if isinstance(node.get('class', []), list) else str(node.get('class', ''))

        # Check callout box
        if 'callout-box' in node_classes_str:
            clean_txt = clean_inline(node)
            if clean_txt:
                elements.append(Paragraph(clean_txt, get_style(callout_style, align_enum)))
            return

        # Check image block or figure
        if node.name in ('figure', 'img') or 'article-image-block' in node_classes_str:
            img_el = node if node.name == 'img' else node.find('img')
            if img_el:
                src = img_el.get('src')
                temp_path = process_image_for_pdf(src)
                if temp_path:
                    img_block = []
                    width_setting = node.get('data-width') or img_el.get('data-width') or ''
                    target_w = 505
                    if width_setting == 'small' or 'img-width-sm' in node_classes_str:
                        target_w = 200
                    elif width_setting == 'medium' or 'img-width-md' in node_classes_str:
                        target_w = 340
                    elif width_setting == 'large' or 'img-width-lg' in node_classes_str:
                        target_w = 420

                    try:
                        pil_i = Image.open(temp_path)
                        w, h = pil_i.size
                        target_h = int((target_w / w) * h)
                        if target_h > 360:
                            target_h = 360
                            target_w = int((target_h / h) * w)
                        img_block.append(RLImage(temp_path, width=target_w, height=target_h))
                    except Exception:
                        img_block.append(Paragraph(f'<img src="{temp_path}" width="{target_w}" />', get_style(body_style, align_enum)))

                    caption_text = None
                    credit_text = None
                    if node.name == 'figure':
                        figcap = node.find('figcaption')
                        if figcap:
                            caption_text = figcap.get_text(strip=True)
                    if not caption_text:
                        em = node.find('em')
                        if em:
                            caption_text = em.get_text(strip=True)
                    if not caption_text and img_el.get('alt'):
                        caption_text = img_el.get('alt')

                    credit_el = node.find(class_=re.compile(r'img-credit|credit'))
                    if credit_el:
                        credit_text = credit_el.get_text(strip=True)

                    if caption_text:
                        img_block.append(Spacer(1, 3))
                        img_block.append(Paragraph(html.escape(caption_text), caption_style))
                    if credit_text:
                        img_block.append(Paragraph(html.escape(credit_text), credit_style))

                    img_block.append(Spacer(1, 8))
                    elements.append(KeepTogether(img_block))
            return

        # Check block tags
        if node.name == 'h1':
            elements.append(Paragraph(clean_inline(node), get_style(h1_style, align_enum)))
        elif node.name == 'h2':
            elements.append(Paragraph(clean_inline(node), get_style(h2_style, align_enum)))
        elif node.name in ('h3', 'h4'):
            elements.append(Paragraph(clean_inline(node), get_style(h3_style, align_enum)))
        elif node.name == 'p':
            nested_imgs = node.find_all('img')
            if nested_imgs:
                for child in node.children:
                    process_node(child)
            else:
                clean_txt = clean_inline(node)
                if clean_txt:
                    elements.append(Paragraph(clean_txt, get_style(body_style, align_enum)))
        elif node.name == 'blockquote':
            clean_txt = clean_inline(node)
            if clean_txt:
                elements.append(Paragraph(f'"{clean_txt}"', get_style(quote_style, align_enum)))
        elif node.name in ('pre', 'code'):
            code_txt = html.escape(node.get_text())
            if code_txt:
                elements.append(Paragraph(code_txt, code_style))
        elif node.name == 'hr':
            elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#E6E2D8'), spaceBefore=10, spaceAfter=10))
        elif node.name in ('ul', 'ol'):
            for li in node.find_all('li', recursive=False):
                clean_li = clean_inline(li)
                if clean_li:
                    elements.append(Paragraph(f"&bull; {clean_li}", get_style(body_style, align_enum)))
        elif node.name in ('div', 'article', 'section', 'main'):
            for child in node.children:
                process_node(child)
        else:
            clean_txt = clean_inline(node)
            if clean_txt:
                elements.append(Paragraph(clean_txt, get_style(body_style, align_enum)))

    for child in root_node.children:
        process_node(child)
            
    # Sources / References
    if article.sources:
        elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#161513'), spaceBefore=20, spaceAfter=10))
        elements.append(Paragraph("Verified Sources & References", ref_head_style))
        
        for idx, source in enumerate(article.sources, 1):
            pub = f"<b>{source.publisher}</b> — " if source.publisher else ""
            ref_text = f"{idx}. {pub}<a href='{source.url}' color='#93602A'><u>{html.escape(source.title)}</u></a>"
            try:
                elements.append(Paragraph(ref_text, ref_body_style))
            except Exception:
                clean_ref = re.sub(r'<[^>]+>', '', ref_text)
                elements.append(Paragraph(clean_ref, ref_body_style))
            
    doc.build(elements, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer


