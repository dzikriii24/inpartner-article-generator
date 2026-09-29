import os
import re
import subprocess
import sys
from datetime import datetime, timedelta

try:
    import openpyxl
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    import openpyxl
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.styles import Font, PatternFill, Alignment

markdown_path = r'c:\Users\dzikri\Downloads\Inpartner\article-generator\project_plan.md'
excel_path = r'c:\Users\dzikri\Downloads\Inpartner\article-generator\project_plan.xlsx'

with open(markdown_path, 'r', encoding='utf-8') as f:
    lines = f.readlines()

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Project Plan"

start_date = datetime(2026, 9, 25)

def parse_days_to_date(day_str):
    # Match "Day 1", "Day 2", etc.
    match = re.search(r'Day\s+(\d+)', day_str, re.IGNORECASE)
    if match:
        day_num = int(match.group(1))
        s_date = start_date + timedelta(days=day_num - 1)
        return s_date.strftime("%d %b %Y")
    return day_str

headers = ['Task', 'PIC', 'Status', 'Date / Due Date', 'Notes']
ws.append(headers)

header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF")
for col_num, cell in enumerate(ws[1], 1):
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')

phase_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
phase_font = Font(bold=True)

current_phase = ""
row_num = 2

for line in lines:
    line = line.strip('\n')
    
    if line.startswith('### '):
        current_phase = line.replace('###', '').strip()
            
        if row_num > 2:
            # Add an empty row for better spacing between phases
            ws.append(['', '', '', '', ''])
            row_num += 1
            
        # Add phase separator row
        ws.append([current_phase, '', '', '', ''])
        ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=5)
        cell = ws.cell(row=row_num, column=1)
        cell.fill = phase_fill
        cell.font = phase_font
        cell.alignment = Alignment(horizontal='left', vertical='center')
        row_num += 1
            
    elif line.startswith('  - ') and not line.startswith('  - **'):
        task = line[4:].strip()
        due_date_formatted = parse_days_to_date(current_phase)
        ws.append([
            task,
            'Dzikri Rabbani', # PIC
            'Belum Berjalan', # Status
            due_date_formatted, # Date / Due Date
            '' # Notes
        ])
        row_num += 1
    elif line.startswith('    - '):
        task = "  - " + line[6:].strip()
        due_date_formatted = parse_days_to_date(current_phase)
        ws.append([
            task,
            'Dzikri Rabbani', # PIC
            'Belum Berjalan', # Status
            due_date_formatted, # Date / Due Date
            '' # Notes
        ])
        row_num += 1

ws.column_dimensions['A'].width = 75
ws.column_dimensions['B'].width = 20
ws.column_dimensions['C'].width = 20
ws.column_dimensions['D'].width = 25
ws.column_dimensions['E'].width = 30

dv = DataValidation(type="list", formula1='"Belum Berjalan,Sedang Berjalan,Sudah Berjalan,Koreksi,Selesai"', allow_blank=True)
dv.error ='Status tidak valid'
dv.errorTitle = 'Pilih dari dropdown'
dv.prompt = 'Pilih status'
dv.promptTitle = 'Status'
ws.add_data_validation(dv)
dv.add(f'C2:C{row_num}')

wb.save(excel_path)
print(f"Excel file created successfully at: {excel_path}")
